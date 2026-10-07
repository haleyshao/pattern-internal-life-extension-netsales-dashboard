#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuilds Pattern's INTERNAL Life Extension Net Sales dashboard HTML from the
same OneDrive ops fill-in workbook used by the external GMV dashboard, and
writes the finished page to OUT_PATH.

Usage:
    python3 build_netsales_dashboard.py <template.html> <out.html>

Unlike the external GMV dashboard's build script, this one DOES read the
净销售 (net sales) block and the two target sheets -- that's the whole
point of this internal-only dashboard. Do not reuse this script's output
anywhere outside Pattern.

The template contains three placeholder lines that this script replaces:
    var BUILD_META = {"generatedAt":"__PENDING__","sourceNote":"__PENDING__"};
    var REAL_RECORDS = [];
    var REAL_TARGETS = {};
"""
import os
import sys
import json
import subprocess
import calendar
import re
import datetime as dt

import openpyxl

SHARE_URL = os.environ.get("ONEDRIVE_URL")

DAILY_SHEET = "每日填报"
TARGET_SHEET = "【无需填写】当月目标及进度"
BYDAY_SHEET = "渠道目标by day拆解"

# 每日填报 sheet: 每日净销售达成 block (date row 33, channel rows 34-43)
NETSALES_DATE_ROW = 33
NETSALES_FIRST_ROW = 34
NETSALES_LAST_ROW = 43
FIRST_DATE_COL = 2  # column B

# channel label (col A) -> subchannel key -- must match SUBCHANNELS in the
# dashboard HTML (same keys as the external GMV dashboard, for consistency).
CHANNEL_MAP = {
    "天猫旗舰店": "tmall_flagship",
    "AliHealth": "tmall_alihealth",
    "TDI": "tmall_tdi",
    "京东自营": "jd_self",
    "京东POP": "jd_pop",
    "抖音": "douyin",
    "小红书": "red",
    "唯品会": "vip",
    "拼多多": "pdd",
    "分销": "b2b",
}

# Haley's B2B/B2C split (2026-10-04): AliHealth, TDI, 京东自营(JD DS) and
# 分销(Sub-dist.) count as B2B; every other channel is B2C. 快手(Kuaishou)
# is intentionally absent -- it's not one of the 10 tracked channels yet.
BIZ_TYPE = {
    "tmall_alihealth": "b2b", "tmall_tdi": "b2b", "jd_self": "b2b", "b2b": "b2b",
    "tmall_flagship": "b2c", "jd_pop": "b2c", "douyin": "b2c", "red": "b2c",
    "vip": "b2c", "pdd": "b2c",
}

DISPLAY_LABELS = {
    "tmall_flagship": "天猫旗舰店", "tmall_alihealth": "AliHealth", "tmall_tdi": "TDI",
    "jd_self": "京东自营", "jd_pop": "京东POP", "douyin": "抖音", "red": "小红书",
    "vip": "唯品会", "pdd": "拼多多", "b2b": "分销",
}

EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def excel_serial_to_iso(serial):
    return (EXCEL_EPOCH + dt.timedelta(days=int(serial))).strftime("%Y-%m-%d")


def cell_to_iso_date(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        return excel_serial_to_iso(value)
    return None


def download_workbook(dest_path):
    if not SHARE_URL:
        raise RuntimeError("ONEDRIVE_URL environment variable is not set")
    url = SHARE_URL + "&download=1"
    cookie_jar = dest_path + ".cookies"
    subprocess.run(
        ["curl", "-sS", "-L", "-b", cookie_jar, "-c", cookie_jar, "-o", dest_path, url],
        check=True,
    )


def parse_netsales_records(ws):
    """Daily 渠道 x date net-sales actuals -- same shape/logic as the
    external dashboard's GMV parser, just pointed at the 净销售达成 block."""
    row_for_channel = {}
    for r in range(NETSALES_FIRST_ROW, NETSALES_LAST_ROW + 1):
        label = ws.cell(row=r, column=1).value
        if label in CHANNEL_MAP:
            row_for_channel[label] = r

    date_cols = []
    c = FIRST_DATE_COL
    empty_run = 0
    max_col = ws.max_column
    while c <= max_col and empty_run <= 10:
        v = ws.cell(row=NETSALES_DATE_ROW, column=c).value
        iso = cell_to_iso_date(v)
        if iso is not None:
            date_cols.append((c, iso))
            empty_run = 0
        elif v is None:
            empty_run += 1
        c += 1

    records = []
    for col, iso_date in date_cols:
        rec = {"date": iso_date}
        any_value = False
        for label, key in CHANNEL_MAP.items():
            row = row_for_channel.get(label)
            val = ws.cell(row=row, column=col).value if row else None
            val = val or 0
            if val:
                any_value = True
            rec[key] = val
        if any_value:
            records.append(rec)

    records.sort(key=lambda r: r["date"])
    return records


def parse_period(val):
    """A block's 统计月份 cell is either a single month (datetime, meaning
    that whole month) or a combined-period string like '2026年9~10月'
    (Haley: Sep had no standalone target since the store opened early, so
    Sep+Oct share one combined target)."""
    if isinstance(val, (dt.datetime, dt.date)):
        y, m = val.year, val.month
        start = dt.date(y, m, 1)
        end = dt.date(y, m, calendar.monthrange(y, m)[1])
        return start, end, f"{y}年{m}月"
    if isinstance(val, str):
        m = re.match(r"(\d{4})年(\d{1,2})~(\d{1,2})月", val)
        if m:
            y, m1, m2 = int(m.group(1)), int(m.group(2)), int(m.group(3))
            start = dt.date(y, m1, 1)
            end = dt.date(y, m2, calendar.monthrange(y, m2)[1])
            return start, end, val
    raise ValueError(f"unrecognized 统计月份 value: {val!r}")


def parse_month_blocks(ws):
    """Scan 【无需填写】当月目标及进度 for every repeating 统计月份 block.
    This sheet is hand-extended by Haley's team one block per month, so we
    must not hardcode how many blocks exist -- discover them all."""
    blocks = []
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "统计月份":
            period_val = ws.cell(row=r, column=4).value
            start, end, label = parse_period(period_val)
            channel_rows = {}
            # scan down until the 合计 row (don't assume exactly 10 channel
            # rows -- ops may add/insert rows, e.g. a new channel)
            for cr in range(r + 3, min(r + 40, ws.max_row + 1)):
                label_cn = ws.cell(row=cr, column=1).value
                if label_cn == "合计" or label_cn == "统计月份":
                    break
                if label_cn in CHANNEL_MAP:
                    channel_rows[CHANNEL_MAP[label_cn]] = cr
            blocks.append({"start": start, "end": end, "label": label, "channel_rows": channel_rows})
    return blocks


def to_num(v):
    """Numbers typed as text ('2,000,000', '¥50000', '5万') still count."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace(",", "").replace("¥", "").replace("元", "").replace(" ", "")
        mult = 1
        if t.endswith("万"):
            t, mult = t[:-1], 10000
        try:
            return float(t) * mult
        except ValueError:
            return None
    return None


SUMMARY_HDR = re.compile(r"(\d{1,2})\s*月\s*(汇总|合计|总计|小计)")


def parse_byday_sections(ws):
    """渠道目标by day拆解 holds several stacked blocks ('GMV目标', '净销售目标').
    For each block return its title and, per channel, the per-day values and
    the monthly 汇总 column values:
        [{"title": str, "channels": {key: {"days": {iso: v}, "sums": {(y, m): v}}}}]
    Block titles / the date-header row are located by content (a col-A title
    containing '目标'; the next row whose col A is '渠道'), not by fixed row
    numbers, so inserted rows or small label edits don't break the read."""
    sections = []
    if ws is None:
        return sections
    titles = []
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if isinstance(label, str) and "目标" in label and label.strip() not in CHANNEL_MAP:
            titles.append((r, label.strip()))
    for r0, title in titles:
        date_row = None
        for r in range(r0, min(r0 + 4, ws.max_row + 1)):
            if ws.cell(row=r, column=1).value == "渠道":
                date_row = r
                break
        if date_row is None:
            date_row = r0 + 1
        date_cols, sum_cols, last_date = [], [], None
        for c in range(2, ws.max_column + 1):
            hv = ws.cell(row=date_row, column=c).value
            iso = cell_to_iso_date(hv)
            if iso:
                date_cols.append((c, iso))
                last_date = iso
            elif isinstance(hv, str):
                m = SUMMARY_HDR.search(hv)
                if m and last_date:
                    sum_cols.append((c, (int(last_date[:4]), int(m.group(1)))))
        channels = {}
        for r in range(date_row + 1, min(date_row + 40, ws.max_row + 1)):
            label = ws.cell(row=r, column=1).value
            if label == "合计":
                break
            if label in CHANNEL_MAP:
                days, sums = {}, {}
                for c, iso in date_cols:
                    v = to_num(ws.cell(row=r, column=c).value)
                    if v:
                        days[iso] = v
                for c, ym in sum_cols:
                    v = to_num(ws.cell(row=r, column=c).value)
                    if v is not None:
                        sums[ym] = v
                channels[CHANNEL_MAP[label]] = {"days": days, "sums": sums}
        sections.append({"title": title, "row": r0, "date_row": date_row,
                         "sum_cols": [ym for _, ym in sum_cols], "channels": channels})
    return sections


def parse_byday_net_targets(ws):
    """Pick the 净销售目标 block out of 渠道目标by day拆解 (never the GMV one) and
    print what every block contained, so the Actions log can be compared with
    Excel line by line."""
    sections = parse_byday_sections(ws)
    for sec in sections:
        tot = {}
        for key, d in sec["channels"].items():
            for ym, v in d["sums"].items():
                tot.setdefault(ym, {})[key] = round(v)
        print(f"[by-day] block '{sec['title']}' (row {sec['row']}, header row {sec['date_row']}) "
              f"汇总 columns={sec['sum_cols']}")
        for ym in sorted(tot):
            nz = {k: v for k, v in tot[ym].items() if v}
            print(f"[by-day]   {ym[0]}-{ym[1]:02d} 汇总 non-zero: {nz or 'none'}")
    net = next((x for x in sections if "净销售" in x["title"]), None)
    if net is None:
        print("[by-day] WARNING: no block with '净销售' in its title was found")
        return {}
    return net["channels"]


def build_targets(ws_target, ws_byday=None):
    """Export every 统计月份 block's per-channel 当月净销售目标.

    Source priority per channel per block:
      1. 渠道目标by day拆解 monthly 汇总 columns (summed over the block's months,
         e.g. 9~10月 block = 9月汇总 + 10月汇总)
      2. col C of 【无需填写】当月目标及进度 (a SUM over those same columns)
      3. sum of the by-day cells inside the block's date range
    The first non-zero wins. A per-block diagnostic is printed ([targets] lines
    in the GitHub Actions log) so what was read can be checked against Excel.

    The dashboard page does all date-dependent math in the browser (so the
    日报 date picker can show any past date).
    """
    byday = parse_byday_net_targets(ws_byday)
    blocks = []
    for b in parse_month_blocks(ws_target):
        s_iso, e_iso = b["start"].strftime("%Y-%m-%d"), b["end"].strftime("%Y-%m-%d")
        months = []
        y, m = b["start"].year, b["start"].month
        while (y, m) <= (b["end"].year, b["end"].month):
            months.append((y, m))
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        targets, notes = {}, []
        for k, r in b["channel_rows"].items():
            d = byday.get(k, {"days": {}, "sums": {}})
            summary_v = sum(d["sums"].get(ym, 0) for ym in months)
            sheet_v = to_num(ws_target.cell(row=r, column=3).value) or 0
            days_v = sum(v for dd, v in d["days"].items() if s_iso <= dd <= e_iso)
            if summary_v:
                chosen, src = summary_v, "汇总列"
            elif sheet_v:
                chosen, src = sheet_v, "当月目标表C列"
            else:
                chosen, src = days_v, "按天格子"
            targets[k] = chosen
            for name, val in (("汇总列", summary_v), ("当月目标表C列", sheet_v), ("按天格子", days_v)):
                if val and name != src and abs(val - chosen) > 1:
                    notes.append(f"{k}: {name}={val:,.0f} differs from chosen {src}={chosen:,.0f}")
        blocks.append({"label": b["label"], "start": s_iso, "end": e_iso, "targets": targets})
        filled = {k: round(v) for k, v in targets.items() if v}
        print(f"[targets] {b['label']} ({s_iso}~{e_iso}) months={months} channels found={len(b['channel_rows'])} "
              f"non-zero net-sales targets={filled or 'NONE'}")
        for n in notes:
            print(f"[targets]   note: {n}")
    blocks.sort(key=lambda x: x["start"])
    return {"blocks": blocks}


def inject(template_path, out_path, records, targets):
    with open(template_path, "r", encoding="utf-8") as f:
        html = f.read()

    marker = (
        'var BUILD_META = {"generatedAt":"__PENDING__","sourceNote":"__PENDING__"};\n'
        '  var REAL_RECORDS = [];\n'
        '  var REAL_TARGETS = {};'
    )
    if marker not in html:
        raise RuntimeError("template marker not found -- did the template change?")

    # GitHub Actions runs in UTC -- stamp the page in Beijing time (UTC+8).
    bj = dt.timezone(dt.timedelta(hours=8))
    generated_at = dt.datetime.now(bj).strftime("%Y-%m-%d %H:%M") + "(北京时间)"
    meta = {
        "generatedAt": generated_at,
        "sourceNote": "运营渠道数据填报表 · 每日填报 · 净销售(内部口径)",
    }
    replacement = (
        "var BUILD_META = " + json.dumps(meta, ensure_ascii=False) + ";\n"
        "  var REAL_RECORDS = " + json.dumps(records, ensure_ascii=False) + ";\n"
        "  var REAL_TARGETS = " + json.dumps(targets, ensure_ascii=False) + ";"
    )
    html = html.replace(marker, replacement)

    full_doc = (
        "<!doctype html>\n<html lang=\"zh-CN\"><head>\n"
        "<meta charset=\"utf-8\">\n"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        "</head><body>\n" + html + "\n</body></html>\n"
    )

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(full_doc)


def main():
    if len(sys.argv) != 3:
        print("usage: build_netsales_dashboard.py <template.html> <out.html>", file=sys.stderr)
        sys.exit(1)
    template_path, out_path = sys.argv[1], sys.argv[2]

    xlsx_tmp = "/tmp/_le_ops_fetch_netsales.xlsx"
    download_workbook(xlsx_tmp)

    wb = openpyxl.load_workbook(xlsx_tmp, data_only=True)
    daily_records = parse_netsales_records(wb[DAILY_SHEET])
    byday = wb[BYDAY_SHEET] if BYDAY_SHEET in wb.sheetnames else None
    targets = build_targets(wb[TARGET_SHEET], byday)

    inject(template_path, out_path, daily_records, targets)
    print(f"wrote {out_path} with {len(daily_records)} day-records, "
          f"{len(targets['blocks'])} target blocks")


if __name__ == "__main__":
    main()
