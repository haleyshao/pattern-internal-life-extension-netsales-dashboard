#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuilds Pattern's INTERNAL Life Extension Net Sales dashboard HTML from the
same OneDrive ops fill-in workbook used by the external GMV dashboard, and
writes the finished page to OUT_PATH.

Usage:
    python3 build_netsales_dashboard.py <template.html> <out.html>

Unlike the external GMV dashboard's build script, this one DOES read the
净销售 (net sales) block and the two target sheets -- that's the whole
point of this internal-only dashboard. Do not reuse this script's output
anywhere outside Pattern.

The template contains three placeholder lines that this script replaces:
    var BUILD_META = {"generatedAt":"__PENDING__","sourceNote":"__PENDING__"};
    var REAL_RECORDS = [];
    var REAL_TARGETS = {};
"""
import os
import sys
import json
import subprocess
import calendar
import re
import datetime as dt

import openpyxl

SHARE_URL = os.environ.get("ONEDRIVE_URL")

DAILY_SHEET = "每日填报"
TARGET_SHEET = "【无需填写】当月目标及进度"

# 每日填报 sheet: 每日净销售达成 block (date row 33, channel rows 34-43)
NETSALES_DATE_ROW = 33
NETSALES_FIRST_ROW = 34
NETSALES_LAST_ROW = 43
FIRST_DATE_COL = 2  # column B

# channel label (col A) -> subchannel key -- must match SUBCHANNELS in the
# dashboard HTML (same keys as the external GMV dashboard, for consistency).
CHANNEL_MAP = {
    "天猫旗舰店": "tmall_flagship",
    "AliHealth": "tmall_alihealth",
    "TDI": "tmall_tdi",
    "京东自营": "jd_self",
    "京东POP": "jd_pop",
    "抖音": "douyin",
    "小红书": "red",
    "唯品会": "vip",
    "拼多多": "pdd",
    "分销": "b2b",
}

# Haley's B2B/B2C split (2026-10-04): AliHealth, TDI, 京东自营(JD DS) and
# 分销(Sub-dist.) count as B2B; every other channel is B2C. 快手(Kuaishou)
# is intentionally absent -- it's not one of the 10 tracked channels yet.
BIZ_TYPE = {
    "tmall_alihealth": "b2b", "tmall_tdi": "b2b", "jd_self": "b2b", "b2b": "b2b",
    "tmall_flagship": "b2c", "jd_pop": "b2c", "douyin": "b2c", "red": "b2c",
    "vip": "b2c", "pdd": "b2c",
}

DISPLAY_LABELS = {
    "tmall_flagship": "天猫旗舰店", "tmall_alihealth": "AliHealth", "tmall_tdi": "TDI",
    "jd_self": "京东自营", "jd_pop": "京东POP", "douyin": "抖音", "red": "小红书",
    "vip": "唯品会", "pdd": "拼多多", "b2b": "分销",
}

EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def excel_serial_to_iso(serial):
    return (EXCEL_EPOCH + dt.timedelta(days=int(serial))).strftime("%Y-%m-%d")


def cell_to_iso_date(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        return excel_serial_to_iso(value)
    return None


def download_workbook(dest_path):
    if not SHARE_URL:
        raise RuntimeError("ONEDRIVE_URL environment variable is not set")
    url = SHARE_URL + "&download=1"
    cookie_jar = dest_path + ".cookies"
    subprocess.run(
        ["curl", "-sS", "-L", "-b", cookie_jar, "-c", cookie_jar, "-o", dest_path, url],
        check=True,
    )


def parse_netsales_records(ws):
    """Daily 渠道 x date net-sales actuals -- same shape/logic as the
    external dashboard's GMV parser, just pointed at the 净销售达成 block."""
    row_for_channel = {}
    for r in range(NETSALES_FIRST_ROW, NETSALES_LAST_ROW + 1):
        label = ws.cell(row=r, column=1).value
        if label in CHANNEL_MAP:
            row_for_channel[label] = r

    date_cols = []
    c = FIRST_DATE_COL
    empty_run = 0
    max_col = ws.max_column
    while c <= max_col and empty_run <= 10:
        v = ws.cell(row=NETSALES_DATE_ROW, column=c).value
        iso = cell_to_iso_date(v)
        if iso is not None:
            date_cols.append((c, iso))
            empty_run = 0
        elif v is None:
            empty_run += 1
        c += 1

    records = []
    for col, iso_date in date_cols:
        rec = {"date": iso_date}
        any_value = False
        for label, key in CHANNEL_MAP.items():
            row = row_for_channel.get(label)
            val = ws.cell(row=row, column=col).value if row else None
            val = val or 0
            if val:
                any_value = True
            rec[key] = val
        if any_value:
            records.append(rec)

    records.sort(key=lambda r: r["date"])
    return records


def parse_period(val):
    """A block's 统计月份 cell is either a single month (datetime, meaning
    that whole month) or a combined-period string like '2026年9~10月'
    (Haley: Sep had no standalone target since the store opened early, so
    Sep+Oct share one combined target)."""
    if isinstance(val, (dt.datetime, dt.date)):
        y, m = val.year, val.month
        start = dt.date(y, m, 1)
        end = dt.date(y, m, calendar.monthrange(y, m)[1])
        return start, end, f"{y}年{m}月"
    if isinstance(val, str):
        m = re.match(r"(\d{4})年(\d{1,2})~(\d{1,2})月", val)
        if m:
            y, m1, m2 = int(m.group(1)), int(m.group(2)), int(m.group(3))
            start = dt.date(y, m1, 1)
            end = dt.date(y, m2, calendar.monthrange(y, m2)[1])
            return start, end, val
    raise ValueError(f"unrecognized 统计月份 value: {val!r}")


def parse_month_blocks(ws):
    """Scan 【无需填写】当月目标及进度 for every repeating 统计月份 block.
    This sheet is hand-extended by Haley's team one block per month, so we
    must not hardcode how many blocks exist -- discover them all."""
    blocks = []
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "统计月份":
            period_val = ws.cell(row=r, column=4).value
            start, end, label = parse_period(period_val)
            channel_rows = {}
            for cr in range(r + 3, r + 13):
                label_cn = ws.cell(row=cr, column=1).value
                if label_cn in CHANNEL_MAP:
                    channel_rows[CHANNEL_MAP[label_cn]] = cr
            blocks.append({"start": start, "end": end, "label": label, "channel_rows": channel_rows})
    return blocks


def group_sum(d, biz):
    return sum(v for k, v in d.items() if BIZ_TYPE.get(k) == biz)


def build_targets(ws_target, daily_records, today):
    blocks = parse_month_blocks(ws_target)

    # YTD actuals: sum every daily net-sales record from Jan 1 of `today`'s
    # calendar year through today (same YTD convention as the external GMV
    # dashboard -- Haley confirmed calendar year, 2026-01-01).
    year_start = dt.date(today.year, 1, 1)
    ytd_by_channel = {}
    for rec in daily_records:
        d = dt.date(*map(int, rec["date"].split("-")))
        if year_start <= d <= today:
            for k, v in rec.items():
                if k == "date":
                    continue
                ytd_by_channel[k] = ytd_by_channel.get(k, 0) + v

    # Annual target: sum 当月净销售目标 (col C) across every block that
    # exists for `today`'s year. Blocks are added by ops one at a time as
    # the year goes on, so this naturally grows from ~0 towards the real
    # full-year number over the course of the year -- that's expected, not
    # a bug (Haley, 2026-10-04: "如果现在抓取是0,就写0").
    annual_target_by_channel = {}
    for b in blocks:
        if b["start"].year != today.year:
            continue
        for k, r in b["channel_rows"].items():
            v = ws_target.cell(row=r, column=3).value or 0
            annual_target_by_channel[k] = annual_target_by_channel.get(k, 0) + v

    def rollup(biz):
        if biz is None:
            target = sum(annual_target_by_channel.values())
            ytd = sum(ytd_by_channel.values())
        else:
            target = group_sum(annual_target_by_channel, biz)
            ytd = group_sum(ytd_by_channel, biz)
        remaining = target - ytd
        rate = (ytd / target) if target else None
        return {"target": target, "ytd": ytd, "remaining": remaining, "rate": rate}

    annual = {"total": rollup(None), "b2b": rollup("b2b"), "b2c": rollup("b2c")}

    # Current-period breakdown table: whichever block's date range contains
    # `today`.
    current_block = next((b for b in blocks if b["start"] <= today <= b["end"]), None)
    current_period = None
    if current_block is not None:
        today_iso = today.strftime("%Y-%m-%d")
        today_rec = next((r for r in daily_records if r["date"] == today_iso), {})
        rows = []
        for k, r in current_block["channel_rows"].items():
            month_target = ws_target.cell(row=r, column=3).value or 0
            mtd = ws_target.cell(row=r, column=5).value or 0
            today_val = today_rec.get(k, 0) or 0
            remaining_target = month_target - mtd
            rate = (mtd / month_target) if month_target else None
            remaining_days = max((current_block["end"] - today).days, 0)
            daily_pace = (remaining_target / remaining_days) if remaining_days > 0 else None
            rows.append({
                "key": k, "label": DISPLAY_LABELS[k], "biz": BIZ_TYPE[k],
                "today": today_val, "target": month_target, "mtd": mtd,
                "remaining": remaining_target, "rate": rate, "pace": daily_pace,
            })

        def subtotal(biz):
            grp = [r for r in rows if biz is None or r["biz"] == biz]
            s_today = sum(r["today"] for r in grp)
            s_target = sum(r["target"] for r in grp)
            s_mtd = sum(r["mtd"] for r in grp)
            s_remaining = sum(r["remaining"] for r in grp)
            s_rate = (s_mtd / s_target) if s_target else None
            remaining_days = max((current_block["end"] - today).days, 0)
            s_pace = (s_remaining / remaining_days) if remaining_days > 0 else None
            return {"today": s_today, "target": s_target, "mtd": s_mtd,
                    "remaining": s_remaining, "rate": s_rate, "pace": s_pace}

        current_period = {
            "label": current_block["label"],
            "end_date": current_block["end"].strftime("%Y-%m-%d"),
            "rows": rows,
            "subtotals": {"b2c": subtotal("b2c"), "b2b": subtotal("b2b"), "total": subtotal(None)},
        }

    return {"annual": annual, "current_period": current_period}


def inject(template_path, out_path, records, targets):
    with open(template_path, "r", encoding="utf-8") as f:
        html = f.read()

    marker = (
        'var BUILD_META = {"generatedAt":"__PENDING__","sourceNote":"__PENDING__"};\n'
        '  var REAL_RECORDS = [];\n'
        '  var REAL_TARGETS = {};'
    )
    if marker not in html:
        raise RuntimeError("template marker not found -- did the template change?")

    generated_at = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = {
        "generatedAt": generated_at,
        "sourceNote": "运营渠道数据填报表 · 每日填报 · 净销售(内部口径)",
    }
    replacement = (
        "var BUILD_META = " + json.dumps(meta, ensure_ascii=False) + ";\n"
        "  var REAL_RECORDS = " + json.dumps(records, ensure_ascii=False) + ";\n"
        "  var REAL_TARGETS = " + json.dumps(targets, ensure_ascii=False) + ";"
    )
    html = html.replace(marker, replacement)

    full_doc = (
        "<!doctype html>\n<html lang=\"zh-CN\"><head>\n"
        "<meta charset=\"utf-8\">\n"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        "</head><body>\n" + html + "\n</body></html>\n"
    )

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(full_doc)


def main():
    if len(sys.argv) != 3:
        print("usage: build_netsales_dashboard.py <template.html> <out.html>", file=sys.stderr)
        sys.exit(1)
    template_path, out_path = sys.argv[1], sys.argv[2]

    xlsx_tmp = "/tmp/_le_ops_fetch_netsales.xlsx"
    download_workbook(xlsx_tmp)

    wb = openpyxl.load_workbook(xlsx_tmp, data_only=True)
    daily_records = parse_netsales_records(wb[DAILY_SHEET])
    today = dt.datetime.now().date()
    targets = build_targets(wb[TARGET_SHEET], daily_records, today)

    inject(template_path, out_path, daily_records, targets)
    print(f"wrote {out_path} with {len(daily_records)} day-records, "
          f"current_period={'yes' if targets['current_period'] else 'NONE (today outside all target blocks)'}")


if __name__ == "__main__":
    main()
